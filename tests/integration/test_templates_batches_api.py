from collections.abc import AsyncIterator, Iterator
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from src.infrastructure.api.main import app
from src.infrastructure.persistence.database import engine
from src.infrastructure.persistence.orm_models import Base

pytestmark = pytest.mark.integration


@pytest.fixture(autouse=True)
async def _clean_database() -> AsyncIterator[None]:
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    yield
    async with engine.begin() as connection:
        for table in reversed(Base.metadata.sorted_tables):
            await connection.execute(table.delete())


@pytest.fixture
def client() -> Iterator[TestClient]:
    with TestClient(app) as test_client:
        yield test_client


def _register(client: TestClient, organization_name: str, email: str) -> dict[str, str]:
    response = client.post(
        "/api/v1/auth/register",
        json={
            "organization_name": organization_name,
            "email": email,
            "password": "supersecret1",
        },
    )
    assert response.status_code == 201, response.text
    return dict(response.json())


def _auth_headers(client: TestClient, organization_name: str, email: str) -> dict[str, str]:
    tokens = _register(client, organization_name, email)
    return {"Authorization": f"Bearer {tokens['access_token']}"}


def _template_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Contrato de {{ arrendatario }}, canon {{ canon_mensual }}.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _excel_bytes(rows: list[dict[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    header = list(dict.fromkeys(key for row in rows for key in row))
    sheet.append(header)
    for row in rows:
        sheet.append([row.get(key, "") for key in header])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload_template(client: TestClient, headers: dict[str, str]) -> dict[str, object]:
    response = client.post(
        "/api/v1/templates",
        headers=headers,
        files={
            "file": (
                "Contrato.docx",
                _template_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    return dict(response.json())


def test_upload_template_detects_placeholders(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")

    template = _upload_template(client, headers)

    assert template["name"] == "Contrato.docx"
    assert set(template["placeholders"]) == {"arrendatario", "canon_mensual"}


def test_list_templates_is_scoped_per_organization(client: TestClient) -> None:
    headers_a = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    headers_b = _auth_headers(client, "Gomez Consultores", "bruno@gomez.co")
    _upload_template(client, headers_a)

    templates_a = client.get("/api/v1/templates", headers=headers_a).json()
    templates_b = client.get("/api/v1/templates", headers=headers_b).json()

    assert len(templates_a) == 1
    assert templates_b == []


def test_submit_batch_creates_pending_batch(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    template = _upload_template(client, headers)
    rows = [
        {"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"},
        {"arrendatario": "Juan Perez", "canon_mensual": "1200000"},
    ]

    response = client.post(
        f"/api/v1/templates/{template['id']}/batches",
        headers=headers,
        files={
            "file": (
                "datos.xlsx",
                _excel_bytes(rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201, response.text
    batch = response.json()
    assert batch["total_rows"] == 2
    assert batch["failed_rows"] == 0
    assert batch["status"] in {"pending", "processing"}


def test_submit_batch_rejects_unknown_template(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    fake_template_id = "00000000-0000-0000-0000-000000000000"

    response = client.post(
        f"/api/v1/templates/{fake_template_id}/batches",
        headers=headers,
        files={"file": ("datos.xlsx", _excel_bytes([{"a": "1"}]), "application/octet-stream")},
    )

    assert response.status_code == 404


def test_get_batch_and_list_documents(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    template = _upload_template(client, headers)
    rows = [{"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"}]
    submitted = client.post(
        f"/api/v1/templates/{template['id']}/batches",
        headers=headers,
        files={"file": ("datos.xlsx", _excel_bytes(rows), "application/octet-stream")},
    ).json()

    batch_response = client.get(f"/api/v1/batches/{submitted['id']}", headers=headers)
    documents_response = client.get(
        f"/api/v1/batches/{submitted['id']}/documents", headers=headers
    )

    assert batch_response.status_code == 200
    assert documents_response.status_code == 200
    documents = documents_response.json()
    assert len(documents) == 1
    assert documents[0]["row_index"] == 0


def test_batches_are_isolated_per_organization(client: TestClient) -> None:
    headers_a = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    headers_b = _auth_headers(client, "Gomez Consultores", "bruno@gomez.co")
    template = _upload_template(client, headers_a)
    excel_content = _excel_bytes([{"arrendatario": "X"}])
    batch = client.post(
        f"/api/v1/templates/{template['id']}/batches",
        headers=headers_a,
        files={"file": ("datos.xlsx", excel_content, "application/octet-stream")},
    ).json()

    response = client.get(f"/api/v1/batches/{batch['id']}", headers=headers_b)

    assert response.status_code == 404


def test_list_batches_without_token_returns_401(client: TestClient) -> None:
    response = client.get("/api/v1/batches")

    assert response.status_code == 401


def test_retry_unknown_batch_returns_404(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    fake_batch_id = "00000000-0000-0000-0000-000000000000"

    response = client.post(f"/api/v1/batches/{fake_batch_id}/retry", headers=headers)

    assert response.status_code == 404


def test_retry_with_nothing_failed_is_a_noop(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    template = _upload_template(client, headers)
    rows = [{"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"}]
    batch = client.post(
        f"/api/v1/templates/{template['id']}/batches",
        headers=headers,
        files={"file": ("datos.xlsx", _excel_bytes(rows), "application/octet-stream")},
    ).json()

    response = client.post(f"/api/v1/batches/{batch['id']}/retry", headers=headers)

    assert response.status_code == 200
    assert response.json()["failed_rows"] == 0


def test_download_unknown_batch_returns_404(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    fake_batch_id = "00000000-0000-0000-0000-000000000000"

    response = client.get(f"/api/v1/batches/{fake_batch_id}/download", headers=headers)

    assert response.status_code == 404


def test_download_batch_returns_a_valid_zip(client: TestClient) -> None:
    headers = _auth_headers(client, "Restrepo & Asociados", "ana@restrepo.co")
    template = _upload_template(client, headers)
    rows = [{"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"}]
    batch = client.post(
        f"/api/v1/templates/{template['id']}/batches",
        headers=headers,
        files={"file": ("datos.xlsx", _excel_bytes(rows), "application/octet-stream")},
    ).json()

    response = client.get(f"/api/v1/batches/{batch['id']}/download", headers=headers)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert f'lote-{batch["id"]}.zip' in response.headers["content-disposition"]
