import asyncio
import subprocess
import sys
import time
from collections.abc import AsyncIterator, Iterator
from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID

import pytest
from docx import Document
from openpyxl import Workbook
from src.application.use_cases.submit_batch import SubmitBatchUseCase
from src.application.use_cases.upload_template import UploadTemplateUseCase
from src.domain.generation.models import BatchStatus, DocumentStatus, GenerationBatch
from src.domain.identity.models import Organization
from src.infrastructure.config import settings
from src.infrastructure.parsing.openpyxl_row_parser import OpenpyxlRowParser
from src.infrastructure.persistence.database import async_session_factory, engine
from src.infrastructure.persistence.orm_models import Base
from src.infrastructure.persistence.postgres_batch_repository import PostgresBatchRepository
from src.infrastructure.persistence.postgres_organization_repository import (
    PostgresOrganizationRepository,
)
from src.infrastructure.persistence.postgres_template_repository import (
    PostgresTemplateRepository,
)
from src.infrastructure.rendering.docxtpl_render_engine import DocxtplRenderEngine
from src.infrastructure.storage.s3_document_storage import S3DocumentStorage
from src.infrastructure.tasks.celery_app import celery_app
from src.infrastructure.tasks.celery_batch_dispatcher import CeleryBatchDispatcher

pytestmark = pytest.mark.integration

_TERMINAL_STATUSES = {BatchStatus.COMPLETED, BatchStatus.FAILED, BatchStatus.COMPLETED_WITH_ERRORS}


@pytest.fixture
async def clean_database() -> AsyncIterator[None]:
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    yield
    async with engine.begin() as connection:
        for table in reversed(Base.metadata.sorted_tables):
            await connection.execute(table.delete())


@pytest.fixture(scope="module")
def celery_worker() -> Iterator[None]:
    process = subprocess.Popen(
        [
            sys.executable,
            "-m",
            "celery",
            "-A",
            "src.infrastructure.tasks.celery_app",
            "worker",
            "--pool=solo",
            "--loglevel=warning",
        ]
    )
    try:
        _wait_until_worker_responds()
        yield
    finally:
        process.terminate()
        process.wait(timeout=10)


def _wait_until_worker_responds(timeout: float = 30.0) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if celery_app.control.ping(timeout=1):
            return
        time.sleep(1)
    raise TimeoutError("El worker de Celery no respondió a tiempo")


def _build_template_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph(
        "Contrato de arrendamiento de {{ arrendatario }}, canon {{ canon_mensual }}."
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_excel_bytes(rows: list[dict[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    header = list(dict.fromkeys(key for row in rows for key in row))
    sheet.append(header)
    for row in rows:
        sheet.append([row.get(key, "") for key in header])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


async def _wait_for_batch(
    organization_id: UUID, batch_id: UUID, timeout: float = 30.0
) -> GenerationBatch:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        async with async_session_factory() as session:
            batch = await PostgresBatchRepository(session).get_by_id(batch_id, organization_id)
        if batch is not None and batch.status in _TERMINAL_STATUSES:
            return batch
        await asyncio.sleep(0.5)
    raise TimeoutError("El lote no terminó de procesarse a tiempo")


async def test_submit_batch_processes_rows_through_celery_and_finalizes(
    clean_database: None, celery_worker: None
) -> None:
    storage = S3DocumentStorage(
        endpoint_url=settings.s3_endpoint_url,
        access_key=settings.s3_access_key,
        secret_key=settings.s3_secret_key,
        bucket=settings.s3_bucket,
    )

    async with async_session_factory() as session:
        organization = Organization.create(
            name="Restrepo & Asociados", created_at=datetime.now(UTC)
        )
        await PostgresOrganizationRepository(session).add(organization)

        upload_use_case = UploadTemplateUseCase(
            PostgresTemplateRepository(session), storage, DocxtplRenderEngine()
        )
        template = await upload_use_case.execute(
            organization.id, "Contrato.docx", _build_template_docx_bytes()
        )

        rows = [
            {"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"},
            {"arrendatario": "Juan Perez", "canon_mensual": "1200000"},
        ]
        submit_use_case = SubmitBatchUseCase(
            PostgresTemplateRepository(session),
            PostgresBatchRepository(session),
            OpenpyxlRowParser(),
            CeleryBatchDispatcher(),
        )
        batch = await submit_use_case.execute(
            organization.id, template.id, _build_excel_bytes(rows)
        )

    final_batch = await _wait_for_batch(organization.id, batch.id)

    assert final_batch.status is BatchStatus.COMPLETED
    assert final_batch.completed_rows == 2
    assert final_batch.failed_rows == 0
    assert final_batch.completed_at is not None

    async with async_session_factory() as session:
        documents = await PostgresBatchRepository(session).list_documents(
            batch.id, organization.id
        )

    assert len(documents) == 2
    assert all(document.status is DocumentStatus.COMPLETED for document in documents)
    for document in documents:
        assert document.output_key is not None
        pdf_bytes = await storage.load(document.output_key)
        assert pdf_bytes.startswith(b"%PDF-")


async def test_submit_batch_with_invalid_rows_completes_with_errors(
    clean_database: None, celery_worker: None
) -> None:
    storage = S3DocumentStorage(
        endpoint_url=settings.s3_endpoint_url,
        access_key=settings.s3_access_key,
        secret_key=settings.s3_secret_key,
        bucket=settings.s3_bucket,
    )

    async with async_session_factory() as session:
        organization = Organization.create(
            name="Gomez Consultores", created_at=datetime.now(UTC)
        )
        await PostgresOrganizationRepository(session).add(organization)

        upload_use_case = UploadTemplateUseCase(
            PostgresTemplateRepository(session), storage, DocxtplRenderEngine()
        )
        template = await upload_use_case.execute(
            organization.id, "Contrato.docx", _build_template_docx_bytes()
        )

        rows = [
            {"arrendatario": "Maria Gonzalez", "canon_mensual": "1500000"},
            {"arrendatario": "Juan Perez"},
        ]
        submit_use_case = SubmitBatchUseCase(
            PostgresTemplateRepository(session),
            PostgresBatchRepository(session),
            OpenpyxlRowParser(),
            CeleryBatchDispatcher(),
        )
        batch = await submit_use_case.execute(
            organization.id, template.id, _build_excel_bytes(rows)
        )

    final_batch = await _wait_for_batch(organization.id, batch.id)

    assert final_batch.status is BatchStatus.COMPLETED_WITH_ERRORS
    assert final_batch.completed_rows == 1
    assert final_batch.failed_rows == 1
