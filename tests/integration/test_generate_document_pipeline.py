from datetime import UTC, datetime
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from docx import Document
from src.application.use_cases.generate_document import GenerateDocumentUseCase
from src.domain.generation.models import (
    DocumentStatus,
    GeneratedDocument,
    GenerationBatch,
    Template,
)
from src.infrastructure.config import settings
from src.infrastructure.pdf.gotenberg_pdf_converter import GotenbergPdfConverter
from src.infrastructure.rendering.docxtpl_render_engine import DocxtplRenderEngine
from src.infrastructure.storage.s3_document_storage import S3DocumentStorage

pytestmark = pytest.mark.integration


def _now() -> datetime:
    return datetime.now(UTC)


def _build_template_docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Contrato de arrendamiento de {{ arrendatario }}.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class FakeTemplateRepository:
    def __init__(self, template: Template) -> None:
        self._template = template

    async def add(self, template: Template) -> None:
        raise NotImplementedError

    async def get_by_id(self, template_id: UUID, organization_id: UUID) -> Template | None:
        if template_id == self._template.id and organization_id == self._template.organization_id:
            return self._template
        return None

    async def list_all(self, organization_id: UUID) -> list[Template]:
        raise NotImplementedError


class FakeBatchRepository:
    def __init__(self, batch: GenerationBatch, document: GeneratedDocument) -> None:
        self._batch = batch
        self._document = document
        self.updated_documents: list[GeneratedDocument] = []

    async def add(self, batch: GenerationBatch, documents: list[GeneratedDocument]) -> None:
        raise NotImplementedError

    async def get_by_id(self, batch_id: UUID, organization_id: UUID) -> GenerationBatch | None:
        if batch_id == self._batch.id and organization_id == self._batch.organization_id:
            return self._batch
        return None

    async def list_all(self, organization_id: UUID) -> list[GenerationBatch]:
        raise NotImplementedError

    async def update(self, batch: GenerationBatch) -> None:
        raise NotImplementedError

    async def get_document(
        self, batch_id: UUID, row_index: int, organization_id: UUID
    ) -> GeneratedDocument | None:
        if batch_id == self._batch.id and row_index == self._document.row_index:
            return self._document
        return None

    async def list_documents(
        self, batch_id: UUID, organization_id: UUID
    ) -> list[GeneratedDocument]:
        raise NotImplementedError

    async def update_document(self, document: GeneratedDocument, organization_id: UUID) -> None:
        self._document = document
        self.updated_documents.append(document)


async def test_pipeline_renders_converts_and_stores_a_real_pdf() -> None:
    organization_id = uuid4()
    storage = S3DocumentStorage(
        endpoint_url=settings.s3_endpoint_url,
        access_key=settings.s3_access_key,
        secret_key=settings.s3_secret_key,
        bucket=settings.s3_bucket,
    )
    template = Template.create(
        organization_id=organization_id,
        name="Contrato_Arrendamiento.docx",
        storage_key=f"templates/{uuid4()}.docx",
        placeholders=["arrendatario"],
        created_at=_now(),
    )
    await storage.save(template.storage_key, _build_template_docx_bytes())

    batch = GenerationBatch.create(
        organization_id=organization_id,
        template_id=template.id,
        total_rows=1,
        created_at=_now(),
    )
    document = GeneratedDocument.create(
        batch_id=batch.id, row_index=0, row_data={"arrendatario": "Maria Gonzalez"}
    )

    batches = FakeBatchRepository(batch, document)
    use_case = GenerateDocumentUseCase(
        template_repository=FakeTemplateRepository(template),
        batch_repository=batches,
        render_engine=DocxtplRenderEngine(),
        pdf_converter=GotenbergPdfConverter(settings.gotenberg_url),
        storage=storage,
    )

    await use_case.execute(organization_id, batch.id, 0)

    final = batches.updated_documents[-1]
    assert final.status is DocumentStatus.COMPLETED
    assert final.output_key is not None

    pdf_bytes = await storage.load(final.output_key)
    assert pdf_bytes.startswith(b"%PDF-")
